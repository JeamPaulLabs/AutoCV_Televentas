# -*- coding: utf-8 -*-
"""Interfaz gráfica para actualizar la información de auditoría (canceladas).

Uso:
    python app.py
    (o doble clic en run.bat en Windows)

Comprueba las dependencias, permite configurar la ubicación del Excel (ruta o URL)
y la conexión a la base de datos, y ejecuta la carga con un botón.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
SETTINGS_FILE = HERE / "settings.json"
REQUIREMENTS = ["openpyxl", "pandas", "psycopg2-binary", "python-dotenv"]

DDL = """
CREATE TABLE IF NOT EXISTS como_vamos_televentas.canceladas_revision (
    periodo date NOT NULL,
    numero_contrato text NOT NULL,
    asesor text,
    producto text,
    fecha_venta date,
    causal_gestion text,
    gestion_supervisor_tmk text,
    observacion text,
    resultado_re_auditoria text,
    observacion_re_auditoria text,
    estado_subsanacion text,
    solicitud text,
    reauditoria text,
    observaciones_reauditoria text,
    fecha_gestion date,
    gestor text,
    tiempo_subs_1 text,
    tiempo_subs_2 text,
    actualizado_en timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (periodo, numero_contrato)
);
"""

MESES_ES = {
    1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO",
    7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE",
    12: "DICIEMBRE",
}

# ---------- Colores (tema Xuma dark) ----------
BG = "#0f172a"
CARD = "#1e293b"
CARD2 = "#273449"
ACCENT = "#0ca30c"
ACCENT_HOVER = "#0fb80f"
TEXT = "#f1f5f9"
MUTED = "#94a3b8"
RED = "#ef4444"
AMBER = "#f59e0b"
GREEN = "#22c55e"


def normalizar_nombre_hoja(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s).strip().upper()


def normalizar_contrato(v):
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:
            return None
        if v == int(v):
            v = int(v)
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "nat"):
        return None
    s = s.replace(" ", "")
    if s.endswith(".0"):
        s = s[:-2]
    return s or None


def normalizar_texto(v):
    if v is None:
        return None
    if isinstance(v, float) and v != v:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "nat"):
        return None
    return s


def parse_fecha(v):
    import pandas as pd
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, (pd.Timestamp,)):
        return None if pd.isna(v) else v.date()
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (pd.to_datetime(v, unit="D", origin="1899-12-30")).date()
        except Exception:
            return None
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "nat"):
        return None
    try:
        ts = pd.to_datetime(s, dayfirst=True, errors="coerce")
        if pd.isna(ts):
            return None
        return ts.date()
    except Exception:
        return None


def norm_header(h):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKD", str(h)).encode("ascii", "ignore").decode("ascii")).strip().upper()


def encontrar_columna(df, *candidatos):
    norm_map = {norm_header(c): c for c in df.columns}
    for cand in candidatos:
        cn = norm_header(cand)
        if cn in norm_map:
            return norm_map[cn]
    return None


def detectar_hoja(wb, periodo, tipo):
    mes = MESES_ES[periodo.month]
    anio = str(periodo.year)[-2:]
    candidatos = []
    for name in wb.sheetnames:
        n = normalizar_nombre_hoja(name)
        if mes not in n or anio not in n:
            continue
        if tipo == "canceladas":
            if "TMK" in n and "SUB" not in n:
                candidatos.append(name)
        else:
            if "SUB" in n and "TMK" in n:
                candidatos.append(name)
    return candidatos[0] if candidatos else None


def leer_canceladas(archivo, periodo, hoja=None):
    import openpyxl
    import pandas as pd
    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    if hoja is None:
        hoja = detectar_hoja(wb, periodo, "canceladas")
    wb.close()
    if hoja is None:
        return pd.DataFrame(), None
    df = pd.read_excel(archivo, sheet_name=hoja)
    col_contrato = encontrar_columna(df, "NUMERO CONTRATO", "NUMEROCONTRATO", "CONTRATO")
    col_asesor = encontrar_columna(df, "ASESOR")
    col_producto = encontrar_columna(df, "PRODUCTO")
    col_fecha_venta = encontrar_columna(df, "FECHA VENTA", "FECHAVENTA")
    col_causal = encontrar_columna(df, "CAUSAL DE GESTION")
    col_gestion = encontrar_columna(df, "GESTION SUPERVISOR TMK", "GESTION  SUPERVISOR TMK")
    col_obs = encontrar_columna(df, "OBSERVACION")
    col_resultado = encontrar_columna(df, "RESULTADO DE RE AUDITORIA Y SUBSANACION")
    col_obs_reaud = encontrar_columna(df, "OBSERVACION RE AUDITORIA Y SUBSANACION")

    if col_contrato is None:
        raise ValueError("No se encontró la columna 'numeroContrato' en la hoja de canceladas.")

    filas = []
    for _, r in df.iterrows():
        contrato = normalizar_contrato(r.get(col_contrato))
        if not contrato:
            continue
        filas.append({
            "periodo": periodo,
            "numero_contrato": contrato,
            "asesor": normalizar_texto(r.get(col_asesor)) if col_asesor else None,
            "producto": normalizar_texto(r.get(col_producto)) if col_producto else None,
            "fecha_venta": parse_fecha(r.get(col_fecha_venta)) if col_fecha_venta else None,
            "causal_gestion": normalizar_texto(r.get(col_causal)) if col_causal else None,
            "gestion_supervisor_tmk": normalizar_texto(r.get(col_gestion)) if col_gestion else None,
            "observacion": normalizar_texto(r.get(col_obs)) if col_obs else None,
            "resultado_re_auditoria": normalizar_texto(r.get(col_resultado)) if col_resultado else None,
            "observacion_re_auditoria": normalizar_texto(r.get(col_obs_reaud)) if col_obs_reaud else None,
        })
    out = pd.DataFrame(filas)
    if not out.empty:
        out = out.drop_duplicates("numero_contrato", keep="last")
    return out, hoja


def leer_subsanacion(archivo, periodo, hoja=None):
    import openpyxl
    import pandas as pd
    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    if hoja is None:
        hoja = detectar_hoja(wb, periodo, "subsanacion")
    wb.close()
    if hoja is None:
        return pd.DataFrame(), None
    df = pd.read_excel(archivo, sheet_name=hoja)
    col_contrato = encontrar_columna(df, "NUMERO CONTRATO", "NUMEROCONTRATO", "CONTRATO")
    col_estado = encontrar_columna(df, "ESTADO DE SUBSANACION")
    col_solicitud = encontrar_columna(df, "SOLICITUD")
    col_reaud = encontrar_columna(df, "REAUDITORIA Y SUBSANACION")
    col_obs = encontrar_columna(df, "OBSERVACIONES")
    col_fecha = encontrar_columna(df, "FECHA GESTION")
    col_gestor = encontrar_columna(df, "GESTOR")
    col_t1 = encontrar_columna(df, "TIEMPO SUBS 1")
    col_t2 = encontrar_columna(df, "TIEMPO SUBS 2")

    if col_contrato is None:
        return pd.DataFrame(), hoja

    filas = []
    for _, r in df.iterrows():
        contrato = normalizar_contrato(r.get(col_contrato))
        if not contrato:
            continue
        filas.append({
            "numero_contrato": contrato,
            "estado_subsanacion": normalizar_texto(r.get(col_estado)) if col_estado else None,
            "solicitud": normalizar_texto(r.get(col_solicitud)) if col_solicitud else None,
            "reauditoria": normalizar_texto(r.get(col_reaud)) if col_reaud else None,
            "observaciones_reauditoria": normalizar_texto(r.get(col_obs)) if col_obs else None,
            "fecha_gestion": parse_fecha(r.get(col_fecha)) if col_fecha else None,
            "gestor": normalizar_texto(r.get(col_gestor)) if col_gestor else None,
            "tiempo_subs_1": normalizar_texto(r.get(col_t1)) if col_t1 else None,
            "tiempo_subs_2": normalizar_texto(r.get(col_t2)) if col_t2 else None,
        })
    out = pd.DataFrame(filas)
    if not out.empty:
        out = out.drop_duplicates("numero_contrato", keep="last")
    return out, hoja


def obtener_archivo(ruta_o_url):
    s = (ruta_o_url or "").strip()
    if s.startswith("http://") or s.startswith("https://"):
        tmp = Path(tempfile.gettempdir()) / "canceladas_gui_archivo.xlsx"
        urllib.request.urlretrieve(s, tmp)
        return tmp
    p = Path(s)
    if not p.exists():
        raise FileNotFoundError(f"No existe el archivo: {p}")
    tmp = Path(tempfile.gettempdir()) / "canceladas_gui_archivo.xlsx"
    shutil.copy2(p, tmp)
    return tmp


def _alcanzable(host, port, timeout=4):
    import socket
    try:
        with socket.create_connection((host, int(port)), timeout=timeout):
            return True
    except (OSError, TypeError, ValueError):
        return False


def _resolver_perfil():
    if _alcanzable("20.7.15.40", 5432):
        return "oficina", "20.7.15.40", 5432
    if _alcanzable("127.0.0.1", 15432):
        return "casa", "127.0.0.1", 15432
    return None, None, None


def conectar_db(cfg):
    import psycopg2
    host = (cfg.get("host") or "").strip()
    port = (cfg.get("port") or "").strip()
    if not host:
        perfil, h, p = _resolver_perfil()
        if h is None:
            raise RuntimeError("No se alcanza ni la oficina (20.7.15.40:5432) ni el túnel (127.0.0.1:15432). "
                               "Si estás en casa, abre el túnel SSH (conectar_tunel.ps1) primero.")
        host, port = h, p
    try:
        return psycopg2.connect(
            host=host, port=int(port), dbname=cfg["database"],
            user=cfg["user"], password=cfg["password"], connect_timeout=8,
        )
    except Exception as e:
        if host == "20.7.15.40" and _alcanzable("127.0.0.1", 15432):
            return psycopg2.connect(
                host="127.0.0.1", port=15432, dbname=cfg["database"],
                user=cfg["user"], password=cfg["password"], connect_timeout=8,
            )
        raise RuntimeError(f"Error de conexión a {host}:{port}: {e}")


def asegurar_tabla(conn):
    with conn.cursor() as cur:
        cur.execute(DDL)
    conn.commit()


def upsert(conn, df):
    import pandas as pd
    if df.empty:
        return 0
    cols = ["periodo", "numero_contrato", "asesor", "producto", "fecha_venta",
            "causal_gestion", "gestion_supervisor_tmk", "observacion",
            "resultado_re_auditoria", "observacion_re_auditoria", "estado_subsanacion",
            "solicitud", "reauditoria", "observaciones_reauditoria", "fecha_gestion",
            "gestor", "tiempo_subs_1", "tiempo_subs_2"]
    placeholders = ", ".join(["%s"] * len(cols))
    colnames = ", ".join(cols)
    update_cols = [c for c in cols if c not in ("periodo", "numero_contrato")]
    update_set = ", ".join(f"{c} = EXCLUDED.{c}" for c in update_cols) + ", actualizado_en = now()"
    query = (f"INSERT INTO como_vamos_televentas.canceladas_revision ({colnames}) "
             f"VALUES ({placeholders}) "
             f"ON CONFLICT (periodo, numero_contrato) DO UPDATE SET {update_set}")
    n = 0
    with conn.cursor() as cur:
        for _, r in df.iterrows():
            vals = []
            for c in cols:
                v = r.get(c)
                if isinstance(v, float) and v != v:
                    v = None
                if isinstance(v, (pd.Timestamp,)):
                    v = None if pd.isna(v) else v.date()
                if c in ("fecha_venta", "fecha_gestion") and not isinstance(v, (dt.date, type(None))):
                    v = None
                vals.append(v)
            cur.execute(query, vals)
            n += 1
        conn.commit()
    return n


def cargar_settings():
    if SETTINGS_FILE.exists():
        try:
            return json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def guardar_settings(cfg):
    SETTINGS_FILE.write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")


def check_deps():
    estado = {}
    estado["python"] = sys.version.split()[0]
    for lib in ["openpyxl", "pandas", "psycopg2", "dotenv"]:
        try:
            m = __import__(lib)
            estado[lib] = getattr(m, "__version__", "ok")
        except Exception:
            estado[lib] = None
    return estado


# ---------- GUI ----------
try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
except Exception as e:
    print("tkinter no disponible:", e)
    sys.exit(1)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Actualizar auditoría — Canceladas")
        self.geometry("780x760")
        self.minsize(700, 600)
        self.configure(bg=BG)
        self.cfg = cargar_settings()
        self._build_styles()
        self._build()
        self._refresh_deps()

    def _build_styles(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure(".", background=BG, foreground=TEXT, font=("Segoe UI", 10))
        style.configure("TFrame", background=BG)
        style.configure("TLabel", background=BG, foreground=TEXT)
        style.configure("TLabelframe", background=CARD, foreground=TEXT, bordercolor=CARD2, relief="flat")
        style.configure("TLabelframe.Label", background=CARD, foreground=ACCENT, font=("Segoe UI", 10, "bold"))
        style.configure("TEntry", fieldbackground=CARD2, foreground=TEXT, bordercolor=CARD2, insertcolor=TEXT)
        style.configure("TButton", background=ACCENT, foreground="#0b1220", font=("Segoe UI", 10, "bold"), padding=6, borderwidth=0)
        style.map("TButton", background=[("active", ACCENT_HOVER), ("disabled", "#475569")])
        style.configure("Secondary.TButton", background=CARD2, foreground=TEXT, padding=6)
        style.map("Secondary.TButton", background=[("active", "#334155")])
        style.configure("TProgressbar", background=ACCENT, troughcolor=CARD2, bordercolor=BG)

    def _build(self):
        # Header
        header = tk.Frame(self, bg=CARD, height=56)
        header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(header, text="  Actualizar auditoría — Canceladas", bg=CARD, fg=TEXT,
                 font=("Segoe UI", 14, "bold")).pack(side="left", pady=12)
        tk.Label(header, text="Xuma · Como Vamos Televentas", bg=CARD, fg=MUTED,
                 font=("Segoe UI", 9)).pack(side="right", padx=12)

        body = tk.Frame(self, bg=BG)
        body.pack(fill="both", expand=True, padx=14, pady=10)

        # --- Dependencias ---
        lf = ttk.LabelFrame(body, text="1. Dependencias")
        lf.pack(fill="x", pady=6)
        self.dep_labels = {}
        for lib in ["python", "openpyxl", "pandas", "psycopg2", "dotenv"]:
            row = tk.Frame(lf, bg=CARD)
            row.pack(fill="x", pady=1)
            tk.Label(row, text=lib, bg=CARD, fg=TEXT, width=14, anchor="w", font=("Segoe UI", 10)).pack(side="left")
            lbl = tk.Label(row, text="...", bg=CARD, fg=MUTED, width=22, anchor="w", font=("Segoe UI", 10))
            lbl.pack(side="left")
            self.dep_labels[lib] = lbl
        ttk.Button(lf, text="Instalar dependencias faltantes", command=self.instalar_deps).pack(anchor="w", pady=6)

        # --- Configuración ---
        cf = ttk.LabelFrame(body, text="2. Configuración")
        cf.pack(fill="x", pady=6)
        self._entry(cf, "Archivo Excel (ruta o URL):", "archivo", width=52)
        ttk.Button(cf, text="Examinar...", command=self.examinar, style="Secondary.TButton").pack(anchor="w", pady=2)
        self._entry(cf, "Mes (YYYY-MM):", "periodo", width=12)
        tk.Label(cf, text="Dejar vacío = mes en curso", bg=CARD, fg=MUTED, font=("Segoe UI", 8)).pack(anchor="w")
        for lbl, key, show in [("Host BD:", "host", None), ("Puerto:", "port", None),
                               ("Base de datos:", "database", None), ("Usuario:", "user", None),
                               ("Contraseña:", "password", "*")]:
            self._entry(cf, lbl, key, width=30, show=show)
        btnrow = tk.Frame(cf, bg=CARD)
        btnrow.pack(fill="x", pady=4)
        ttk.Button(btnrow, text="Auto-detectar perfil", command=self.auto_detectar, style="Secondary.TButton").pack(side="left", padx=2)
        ttk.Button(btnrow, text="Probar conexión", command=self.probar_conexion, style="Secondary.TButton").pack(side="left", padx=2)
        ttk.Button(btnrow, text="Guardar configuración", command=self.guardar).pack(side="left", padx=2)

        # --- Acción ---
        af = ttk.LabelFrame(body, text="3. Ejecutar")
        af.pack(fill="x", pady=6)
        self.progress = ttk.Progressbar(af, mode="indeterminate", length=200)
        self.progress.pack(fill="x", pady=4)
        ttk.Button(af, text="Actualizar auditoría (cargar canceladas)", command=self.ejecutar).pack(anchor="w", pady=4)

        # --- Log ---
        lf2 = ttk.LabelFrame(body, text="Log")
        lf2.pack(fill="both", expand=True, pady=6)
        self.log = tk.Text(lf2, height=14, wrap="word", bg="#0b1220", fg="#22c55e",
                           font=("Consolas", 9), insertbackground=TEXT, relief="flat", padx=8, pady=6)
        self.log.pack(fill="both", expand=True)

        # Footer
        footer = tk.Frame(self, bg=CARD, height=28)
        footer.pack(fill="x", side="bottom")
        footer.pack_propagate(False)
        tk.Label(footer, text="settings.json (local, no se sube a git) · Sin datos personales del asegurado",
                 bg=CARD, fg=MUTED, font=("Segoe UI", 8)).pack(side="left", padx=12, pady=6)

    def _entry(self, parent, label, key, width=40, show=None):
        f = tk.Frame(parent, bg=CARD)
        f.pack(fill="x", pady=2)
        tk.Label(f, text=label, bg=CARD, fg=TEXT, width=26, anchor="w", font=("Segoe UI", 10)).pack(side="left")
        e = tk.Entry(f, width=width, show=show or "", bg=CARD2, fg=TEXT, insertbackground=TEXT,
                     relief="flat", highlightthickness=1, highlightbackground=CARD2, highlightcolor=ACCENT,
                     font=("Segoe UI", 10))
        e.pack(side="left", fill="x", expand=True, ipady=3)
        val = self.cfg.get(key, "")
        if key == "periodo" and not val:
            val = dt.date.today().strftime("%Y-%m")
        if key == "port" and not val:
            val = "5432"
        if key == "host" and not val:
            val = "20.7.15.40"
        if key == "database" and not val:
            val = "DataCenter_Vanti"
        e.insert(0, str(val))
        setattr(self, f"e_{key}", e)

    def _get(self, key):
        return getattr(self, f"e_{key}").get().strip()

    def _log(self, msg):
        self.log.insert("end", msg + "\n")
        self.log.see("end")
        self.update_idletasks()

    def _refresh_deps(self):
        st = check_deps()
        self.dep_labels["python"].config(text=st["python"], fg=GREEN)
        for lib in ["openpyxl", "pandas", "psycopg2", "dotenv"]:
            v = st[lib]
            if v:
                self.dep_labels[lib].config(text=v, fg=GREEN)
            else:
                self.dep_labels[lib].config(text="FALTA", fg=RED)
        self._log("Dependencias: " + ", ".join(f"{k}={v}" for k, v in st.items()))

    def instalar_deps(self):
        self._log("Instalando dependencias...")
        self.progress.start(10)
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "-r", str(HERE / "requirements.txt")])
            self._log("Dependencias instaladas.")
        except Exception as e:
            self._log("Error instalando: " + str(e))
            messagebox.showerror("Error", str(e))
        finally:
            self.progress.stop()
        self._refresh_deps()

    def examinar(self):
        p = filedialog.askopenfilename(
            title="Seleccionar Excel de canceladas",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Todos", "*.*")],
        )
        if p:
            self.e_archivo.delete(0, "end")
            self.e_archivo.insert(0, p)

    def guardar(self):
        cfg = {
            "archivo": self._get("archivo"),
            "periodo": self._get("periodo"),
            "host": self._get("host"),
            "port": self._get("port"),
            "database": self._get("database"),
            "user": self._get("user"),
            "password": self._get("password"),
        }
        guardar_settings(cfg)
        self.cfg = cfg
        self._log("Configuración guardada.")

    def auto_detectar(self):
        perfil, host, port = _resolver_perfil()
        if host is None:
            messagebox.showwarning("Perfil", "No se alcanza ni la oficina (20.7.15.40:5432) ni el túnel (127.0.0.1:15432).\n"
                                            "Si estás en casa, abre el túnel SSH (conectar_tunel.ps1) primero.")
            self._log("Auto-detectar: sin conexión disponible.")
            return
        self.e_host.delete(0, "end"); self.e_host.insert(0, host)
        self.e_port.delete(0, "end"); self.e_port.insert(0, str(port))
        self._log(f"Auto-detectar: perfil {perfil} -> {host}:{port}")
        messagebox.showinfo("Perfil", f"Perfil detectado: {perfil} ({host}:{port})")

    def probar_conexion(self):
        cfg = self._cfg_actual()
        self.progress.start(10)
        try:
            conn = conectar_db(cfg)
            conn.close()
            self._log("Conexión OK.")
            messagebox.showinfo("Conexión", "Conexión exitosa.")
        except Exception as e:
            self._log("Error de conexión: " + str(e))
            messagebox.showerror("Conexión", str(e))
        finally:
            self.progress.stop()

    def _cfg_actual(self):
        return {
            "archivo": self._get("archivo"),
            "periodo": self._get("periodo"),
            "host": self._get("host"),
            "port": self._get("port"),
            "database": self._get("database"),
            "user": self._get("user"),
            "password": self._get("password"),
        }

    def ejecutar(self):
        cfg = self._cfg_actual()
        self.guardar()
        self._log("=== Iniciando carga ===")
        self.progress.start(10)
        try:
            periodo_txt = cfg["periodo"] or dt.date.today().strftime("%Y-%m")
            periodo = dt.datetime.strptime(periodo_txt, "%Y-%m").date().replace(day=1)
            archivo = obtener_archivo(cfg["archivo"])
            self._log(f"Periodo: {periodo:%Y-%m}")
            self._log(f"Archivo: {cfg['archivo']}")

            conn = conectar_db(cfg)
            asegurar_tabla(conn)

            df_canc, hoja_c = leer_canceladas(archivo, periodo)
            self._log(f"Hoja canceladas: {hoja_c} -> {len(df_canc)} válidas")
            df_sub, hoja_s = leer_subsanacion(archivo, periodo)
            self._log(f"Hoja subsanación: {hoja_s} -> {len(df_sub)} válidas")

            if df_canc.empty:
                self._log("Sin canceladas para cargar. Fin.")
                conn.close()
                return

            if not df_sub.empty:
                df = df_canc.merge(df_sub, on="numero_contrato", how="left", suffixes=("", "_sub"))
                for col in ("estado_subsanacion", "solicitud", "reauditoria",
                            "observaciones_reauditoria", "fecha_gestion", "gestor",
                            "tiempo_subs_1", "tiempo_subs_2"):
                    alt = col + "_sub"
                    if alt in df.columns:
                        df[col] = df[alt].combine_first(df.get(col))
                        df = df.drop(columns=[alt])
            else:
                df = df_canc.copy()
                for col in ("estado_subsanacion", "solicitud", "reauditoria",
                            "observaciones_reauditoria", "fecha_gestion", "gestor",
                            "tiempo_subs_1", "tiempo_subs_2"):
                    df[col] = None

            df["periodo"] = periodo
            cols = ["periodo", "numero_contrato", "asesor", "producto", "fecha_venta",
                    "causal_gestion", "gestion_supervisor_tmk", "observacion",
                    "resultado_re_auditoria", "observacion_re_auditoria", "estado_subsanacion",
                    "solicitud", "reauditoria", "observaciones_reauditoria", "fecha_gestion",
                    "gestor", "tiempo_subs_1", "tiempo_subs_2"]
            for c in cols:
                if c not in df.columns:
                    df[c] = None
            df = df[cols]

            self._log(f"Filas a escribir: {len(df)}")
            n = upsert(conn, df)
            conn.close()
            self._log(f"Filas escritas/actualizadas: {n}")
            self._log("OK")
            messagebox.showinfo("Listo", f"Carga completada: {n} filas.")
        except Exception as e:
            self._log("ERROR: " + str(e))
            messagebox.showerror("Error", str(e))
        finally:
            self.progress.stop()


if __name__ == "__main__":
    app = App()
    app.mainloop()
