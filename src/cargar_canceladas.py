# -*- coding: utf-8 -*-
"""Carga la revisión (subsanación / re-auditoría) de las ventas canceladas
desde "REPORTE CANCELADAS 2026 CMK VANTI.xlsx" hacia la tabla
como_vamos_televentas.canceladas_revision.

Fuente: hojas "TMK <MES><AÑO>" (canceladas) y "SUB TMK <MES> <AÑO>"
(re-auditoría). No se cargan datos personales del asegurado.

Uso:
    python src/cargar_canceladas.py                      # mes en curso, archivo por defecto
    python src/cargar_canceladas.py --periodo 2026-10
    python src/cargar_canceladas.py --archivo "C:\\ruta\\al\\archivo.xlsx"
    python src/cargar_canceladas.py --hoja-canceladas "TMK OCTUBRE26" --hoja-subsanacion "SUB TMK OCTUBRE 26"
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.db import get_conexion

ARCHIVO_DEFAULT = (
    r"C:\Users\jarcon\OneDrive - Chariot & Castle Seguros SAS\Shortcuts"
    r"\POSVENTA - XUMA - REPORTES DIARIOS - VANTI\REPORTE CANCELADAS 2026 CMK VANTI.xlsx"
)

MESES_ES = {
    1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO",
    7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE",
    12: "DICIEMBRE",
}

DDL_FALLBACK = """
CREATE TABLE IF NOT EXISTS como_vamos_televentas.canceladas_revision (
    periodo date NOT NULL,
    numero_contrato text NOT NULL,
    asesor text,
    producto text,
    fecha_venta date,
    causal_gestion text,
    gestion_supervisor_tmk text,
    observacion text,
    resultado_re_auditoria text,
    observacion_re_auditoria text,
    estado_subsanacion text,
    solicitud text,
    reauditoria text,
    observaciones_reauditoria text,
    fecha_gestion date,
    gestor text,
    tiempo_subs_1 text,
    tiempo_subs_2 text,
    actualizado_en timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (periodo, numero_contrato)
);
"""


def normalizar_nombre_hoja(s) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s).strip().upper()


def normalizar_contrato(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        if v == int(v):
            v = int(v)
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "nat"):
        return None
    s = s.replace(" ", "")
    if s.endswith(".0"):
        s = s[:-2]
    return s or None


def normalizar_texto(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v != v:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "nat"):
        return None
    return s


def parse_fecha(v) -> dt.date | None:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, (pd.Timestamp,)):
        return None if pd.isna(v) else v.date()
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (pd.to_datetime(v, unit="D", origin="1899-12-30")).date()
        except Exception:
            return None
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "nat"):
        return None
    try:
        ts = pd.to_datetime(s, dayfirst=True, errors="coerce")
        if pd.isna(ts):
            return None
        return ts.date()
    except Exception:
        return None


def norm_header(h) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKD", str(h)).encode("ascii", "ignore").decode("ascii")).strip().upper()


def encontrar_columna(df, *candidatos) -> str | None:
    norm_map = {norm_header(c): c for c in df.columns}
    for cand in candidatos:
        cn = norm_header(cand)
        if cn in norm_map:
            return norm_map[cn]
    return None


def detectar_hoja(wb, periodo, tipo) -> str | None:
    mes = MESES_ES[periodo.month]
    anio = str(periodo.year)[-2:]
    candidatos = []
    for name in wb.sheetnames:
        n = normalizar_nombre_hoja(name)
        if mes not in n or anio not in n:
            continue
        if tipo == "canceladas":
            if "TMK" in n and "SUB" not in n:
                candidatos.append(name)
        else:  # subsanacion
            if "SUB" in n and "TMK" in n:
                candidatos.append(name)
    return candidatos[0] if candidatos else None


def leer_canceladas(archivo: Path, periodo: dt.date, hoja: str | None) -> pd.DataFrame:
    if hoja is None:
        import openpyxl
        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        hoja = detectar_hoja(wb, periodo, "canceladas")
        wb.close()
    if hoja is None:
        print(f"AVISO: no se encontró hoja de canceladas TMK para {periodo:%Y-%m}.")
        return pd.DataFrame()

    df = pd.read_excel(archivo, sheet_name=hoja)
    print(f"Hoja canceladas: '{hoja}' -> {len(df)} filas leídas.")

    col_contrato = encontrar_columna(df, "NUMERO CONTRATO", "NUMEROCONTRATO", "CONTRATO")
    col_asesor = encontrar_columna(df, "ASESOR")
    col_producto = encontrar_columna(df, "PRODUCTO")
    col_fecha_venta = encontrar_columna(df, "FECHA VENTA", "FECHAVENTA")
    col_causal = encontrar_columna(df, "CAUSAL DE GESTION")
    col_gestion = encontrar_columna(df, "GESTION SUPERVISOR TMK", "GESTION  SUPERVISOR TMK")
    col_obs = encontrar_columna(df, "OBSERVACION")
    col_resultado = encontrar_columna(df, "RESULTADO DE RE AUDITORIA Y SUBSANACION")
    col_obs_reaud = encontrar_columna(df, "OBSERVACION RE AUDITORIA Y SUBSANACION")

    if col_contrato is None:
        raise SystemExit("No se encontró la columna 'numeroContrato' en la hoja de canceladas.")

    filas = []
    for _, r in df.iterrows():
        contrato = normalizar_contrato(r.get(col_contrato) if col_contrato else None)
        if not contrato:
            continue
        filas.append({
            "periodo": periodo,
            "numero_contrato": contrato,
            "asesor": normalizar_texto(r.get(col_asesor)) if col_asesor else None,
            "producto": normalizar_texto(r.get(col_producto)) if col_producto else None,
            "fecha_venta": parse_fecha(r.get(col_fecha_venta)) if col_fecha_venta else None,
            "causal_gestion": normalizar_texto(r.get(col_causal)) if col_causal else None,
            "gestion_supervisor_tmk": normalizar_texto(r.get(col_gestion)) if col_gestion else None,
            "observacion": normalizar_texto(r.get(col_obs)) if col_obs else None,
            "resultado_re_auditoria": normalizar_texto(r.get(col_resultado)) if col_resultado else None,
            "observacion_re_auditoria": normalizar_texto(r.get(col_obs_reaud)) if col_obs_reaud else None,
        })
    out = pd.DataFrame(filas)
    # dedupe por contrato (último gana) dentro del mismo mes
    if not out.empty:
        out = out.drop_duplicates("numero_contrato", keep="last")
    print(f"  canceladas válidas (contrato no vacío): {len(out)}")
    return out


def leer_subsanacion(archivo: Path, periodo: dt.date, hoja: str | None) -> pd.DataFrame:
    if hoja is None:
        import openpyxl
        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        hoja = detectar_hoja(wb, periodo, "subsanacion")
        wb.close()
    if hoja is None:
        print(f"AVISO: no se encontró hoja SUB TMK para {periodo:%Y-%m} (re-auditoría).")
        return pd.DataFrame()

    df = pd.read_excel(archivo, sheet_name=hoja)
    print(f"Hoja subsanación: '{hoja}' -> {len(df)} filas leídas.")

    col_contrato = encontrar_columna(df, "NUMERO CONTRATO", "NUMEROCONTRATO", "CONTRATO")
    col_estado = encontrar_columna(df, "ESTADO DE SUBSANACION")
    col_solicitud = encontrar_columna(df, "SOLICITUD")
    col_reaud = encontrar_columna(df, "REAUDITORIA Y SUBSANACION")
    col_obs = encontrar_columna(df, "OBSERVACIONES")
    col_fecha = encontrar_columna(df, "FECHA GESTION")
    col_gestor = encontrar_columna(df, "GESTOR")
    col_t1 = encontrar_columna(df, "TIEMPO SUBS 1")
    col_t2 = encontrar_columna(df, "TIEMPO SUBS 2")

    if col_contrato is None:
        print("AVISO: sin columna de contrato en la hoja de subsanación.")
        return pd.DataFrame()

    filas = []
    for _, r in df.iterrows():
        contrato = normalizar_contrato(r.get(col_contrato))
        if not contrato:
            continue
        filas.append({
            "numero_contrato": contrato,
            "estado_subsanacion": normalizar_texto(r.get(col_estado)) if col_estado else None,
            "solicitud": normalizar_texto(r.get(col_solicitud)) if col_solicitud else None,
            "reauditoria": normalizar_texto(r.get(col_reaud)) if col_reaud else None,
            "observaciones_reauditoria": normalizar_texto(r.get(col_obs)) if col_obs else None,
            "fecha_gestion": parse_fecha(r.get(col_fecha)) if col_fecha else None,
            "gestor": normalizar_texto(r.get(col_gestor)) if col_gestor else None,
            "tiempo_subs_1": normalizar_texto(r.get(col_t1)) if col_t1 else None,
            "tiempo_subs_2": normalizar_texto(r.get(col_t2)) if col_t2 else None,
        })
    out = pd.DataFrame(filas)
    if not out.empty:
        out = out.drop_duplicates("numero_contrato", keep="last")
    print(f"  subsanación válidas: {len(out)}")
    return out


def asegurar_tabla():
    ddl_path = ROOT / "sql" / "ddl_canceladas_revision.sql"
    sql = ddl_path.read_text(encoding="utf-8") if ddl_path.exists() else DDL_FALLBACK
    with get_conexion() as conn:
        cur = conn.cursor()
        cur.execute(sql)
        conn.commit()


def upsert(df: pd.DataFrame) -> int:
    if df.empty:
        return 0
    cols = ["periodo", "numero_contrato", "asesor", "producto", "fecha_venta",
            "causal_gestion", "gestion_supervisor_tmk", "observacion",
            "resultado_re_auditoria", "observacion_re_auditoria", "estado_subsanacion",
            "solicitud", "reauditoria", "observaciones_reauditoria", "fecha_gestion",
            "gestor", "tiempo_subs_1", "tiempo_subs_2"]
    placeholders = ", ".join(["%s"] * len(cols))
    colnames = ", ".join(cols)
    update_cols = [c for c in cols if c not in ("periodo", "numero_contrato")]
    update_set = ", ".join(f"{c} = EXCLUDED.{c}" for c in update_cols) + ", actualizado_en = now()"
    query = (
        f"INSERT INTO como_vamos_televentas.canceladas_revision ({colnames}) "
        f"VALUES ({placeholders}) "
        f"ON CONFLICT (periodo, numero_contrato) DO UPDATE SET {update_set}"
    )
    n = 0
    with get_conexion() as conn:
        cur = conn.cursor()
        for _, r in df.iterrows():
            vals = []
            for c in cols:
                v = r.get(c)
                if isinstance(v, float) and v != v:
                    v = None
                if isinstance(v, (pd.Timestamp,)):
                    v = None if pd.isna(v) else v.date()
                if c in ("fecha_venta", "fecha_gestion") and not isinstance(v, (dt.date, type(None))):
                    v = None
                vals.append(v)
            cur.execute(query, vals)
            n += 1
        conn.commit()
    return n


def main():
    ap = argparse.ArgumentParser(description="Carga canceladas + revisión (subsanación/re-auditoría) a la BD.")
    ap.add_argument("--periodo", default=None, help="YYYY-MM (default: mes en curso)")
    ap.add_argument("--archivo", default=ARCHIVO_DEFAULT, help="Ruta al Excel")
    ap.add_argument("--hoja-canceladas", default=None, help="Nombre exacto de la hoja de canceladas TMK")
    ap.add_argument("--hoja-subsanacion", default=None, help="Nombre exacto de la hoja SUB TMK")
    args = ap.parse_args()

    periodo = (
        dt.datetime.strptime(args.periodo, "%Y-%m").date().replace(day=1)
        if args.periodo else dt.date.today().replace(day=1)
    )
    archivo = Path(args.archivo)
    if not archivo.exists():
        raise SystemExit(f"No se encontró el archivo: {archivo}")

    # Copiar a una ruta temporal antes de leer: el archivo suele estar abierto en
    # Excel (bloqueado) y pd.read_excel fallaría con PermissionError. La copia
    # se puede leer aunque el original esté bloqueado.
    import shutil
    import tempfile
    tmp_dir = Path(tempfile.gettempdir()) / "cargar_canceladas"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    archivo_tmp = tmp_dir / archivo.name
    try:
        shutil.copy2(archivo, archivo_tmp)
    except Exception as e:
        raise SystemExit(f"No se pudo copiar el archivo (¿está abierto en Excel?): {e}")
    archivo = archivo_tmp
    print(f"Periodo: {periodo:%Y-%m}")
    print(f"Archivo: {args.archivo}")
    print(f"Copia temporal: {archivo_tmp}")

    asegurar_tabla()

    df_canc = leer_canceladas(archivo, periodo, args.hoja_canceladas)
    df_sub = leer_subsanacion(archivo, periodo, args.hoja_subsanacion)

    if df_canc.empty:
        print("Sin canceladas para cargar. Fin.")
        return

    # Merge por contrato (re-auditoría sobrescribe/añade campos)
    if not df_sub.empty:
        df = df_canc.merge(df_sub, on="numero_contrato", how="left", suffixes=("", "_sub"))
        # si hubo columnas duplicadas por el merge, quedarse con la de subsanación
        for col in ("estado_subsanacion", "solicitud", "reauditoria",
                    "observaciones_reauditoria", "fecha_gestion", "gestor",
                    "tiempo_subs_1", "tiempo_subs_2"):
            alt = col + "_sub"
            if alt in df.columns:
                df[col] = df[alt].combine_first(df.get(col))
                df = df.drop(columns=[alt])
    else:
        df = df_canc.copy()
        for col in ("estado_subsanacion", "solicitud", "reauditoria",
                    "observaciones_reauditoria", "fecha_gestion", "gestor",
                    "tiempo_subs_1", "tiempo_subs_2"):
            df[col] = None

    df["periodo"] = periodo
    # orden de columnas
    cols = ["periodo", "numero_contrato", "asesor", "producto", "fecha_venta",
            "causal_gestion", "gestion_supervisor_tmk", "observacion",
            "resultado_re_auditoria", "observacion_re_auditoria", "estado_subsanacion",
            "solicitud", "reauditoria", "observaciones_reauditoria", "fecha_gestion",
            "gestor", "tiempo_subs_1", "tiempo_subs_2"]
    for c in cols:
        if c not in df.columns:
            df[c] = None
    df = df[cols]

    print(f"Filas a escribir: {len(df)}")
    n = upsert(df)
    print(f"Filas escritas/actualizadas: {n}")
    print("OK")


if __name__ == "__main__":
    main()
